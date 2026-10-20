"""Vigia da inclusão — PROCESSO_INCLUSAO_HAPVIDA.md (fechado com o André em 15/10/2026), passos 13 e 17–19.

Cada inclusão feita pela Cora fica em saneamento.inclusao até cada vida estar ativa ou encerrada. Roda pelo cron
(a cada 30 min, horário comercial) e faz, por inclusão aberta:
  - confere na Hapvida quem já ficou ativo (no máximo a cada 2h) e grava essas vidas no nosso banco (passo 13);
  - na mesma conferência, pedido que a Hapvida CANCELOU (205/206 sem documento) roda o motor de novo e sobe de novo
    pra pendenciar, sem aprovação (19/10/2026, André: "inclui, deixa pendenciar, relatório");
  - todas ativas -> card final (passo 19) + o relatório de 49 colunas do /fechar-contrato só com as vidas da inclusão;
  - quando nada mais está em andamento no contrato (o motor esgotou) e ainda falta vida, pergunta ao pós-vendas:
    Continuar incluindo / Finalizar (passo 17);
  - Continuar: lembrete a cada 8h úteis, com o lote/protocolo e "Quero o arquivo";
  - prazo de 30 dias do evento (17/10/2026): gestor em cópia 10 dias antes; só finaliza depois do prazo vencer
    (sem data de evento no arquivo: 4 dias úteis sem resposta -> gestor; 7 -> finaliza).
A decisão do cartão chega por app.py /inclusao/decisao -> decidir().

Uso: python3 ciclo_inclusao.py [--simular]   (--simular: só mostra o que faria, não manda nada nem grava)
"""
import base64
import io
import json
import os
import sys
import threading
from datetime import datetime, timedelta
from pathlib import Path

import psycopg2
import psycopg2.extras
import requests
from openpyxl import load_workbook

sys.path.insert(0, os.environ.get("SANEAMENTO_PATH", "/lifetools_service"))
from saneamento import coleta_portal  # noqa: E402
from saneamento.regras_inclusao import prazo_evento  # noqa: E402
from saneamento.texto import normalizar  # noqa: E402

CORA = os.environ.get("CORA_BASE_URL", "http://cora:3978")
API = os.environ.get("CONVERTER_ENDPOINT", "http://lifetools-api-1:8000")
ADMIN = os.environ.get("ADMIN_UPN", "andre.aguiar@ra-beneficios.com.br")
SAIDAS = Path(os.environ.get("SAIDA_MOTOR", "/dados_motor/_saidas"))
CONFERIR_A_CADA = timedelta(hours=2)
PRIMEIRA_CONFERENCIA = timedelta(hours=1)
LEMBRETE_HORAS_UTEIS = 8
DIAS_GESTOR, DIAS_FINALIZA = 4, 7   # sem data de evento no arquivo (fallback)
AVISO_GESTOR_ANTES_DO_PRAZO = 10    # 17/10/2026 (André): gestor avisado 10 dias antes do fim do prazo de 30 dias do evento
EXPEDIENTE = (8, 18)            # horário comercial: seg–sex, 8h–18h
PLANO_ODONTO = "9732"           # 117 do Odonto (regra do André, 24/09)
FINAIS_ACOMPANHAMENTO = ("CONCLUIDO", "FALHOU")
DADOS = Path(os.environ.get("CICLO_DADOS", "/dados_motor"))
STATUS_FINAIS = {"FileProcessed", "FileProcessedWithReview", "ProcessFileWithPendency"}


class Recusa(Exception):
    pass


def _conn():
    return psycopg2.connect(host=os.environ.get("DB_HOST", "postgres"), user=os.environ.get("DB_USER", "radb"),
                            password=os.environ.get("DB_PASSWORD", ""), dbname=os.environ.get("DB_NAME", "lifetools"))


def _cora(rota: str, corpo: dict) -> bool:
    try:
        r = requests.post(f"{CORA}{rota}", json=corpo, timeout=120)
        return r.ok and (r.json() or {}).get("enviado", True) is not False
    except Exception as e:
        print(f"[inclusao] Cora {rota} falhou: {e}")
        return False


def horas_uteis(de: datetime, ate: datetime) -> float:
    """Horas de expediente (seg–sex, 8h–18h) entre dois instantes (naive, horário local)."""
    total, t = 0.0, de
    while t < ate:
        fim_dia = t.replace(hour=EXPEDIENTE[1], minute=0, second=0, microsecond=0)
        ini_dia = t.replace(hour=EXPEDIENTE[0], minute=0, second=0, microsecond=0)
        if t.weekday() < 5:
            a, b = max(t, ini_dia), min(ate, fim_dia)
            if b > a:
                total += (b - a).total_seconds() / 3600
        t = (t + timedelta(days=1)).replace(hour=0, minute=0, second=0, microsecond=0)
    return total


def no_expediente(agora: datetime) -> bool:
    return agora.weekday() < 5 and EXPEDIENTE[0] <= agora.hour < EXPEDIENTE[1]


def _gestor(cur) -> str | None:
    cur.execute("SELECT upn FROM saneamento.papel WHERE papel = 'GESTOR_POS_VENDAS'")
    row = cur.fetchone()
    return row["upn"] if row else None


def _naive(d):
    return d.replace(tzinfo=None) if d and d.tzinfo else d


# ── planilha da inclusão (layout de 49 colunas, só as linhas que subiram) ─────────────────────────────────
def _linhas_xlsx(xlsx: bytes) -> tuple[list, list[list], int, dict]:
    """(cabeçalho, linhas, índice da coluna, mapa código->coluna)."""
    ws = load_workbook(io.BytesIO(xlsx)).active
    cab = [c.value for c in ws[1]]
    idx = {}
    for n, h in enumerate(cab):
        cod = str(h or "").split(" ")[0].strip().lower()
        if cod:
            idx[cod] = n
    linhas = [[c.value for c in r] for r in ws.iter_rows(min_row=2) if any(c.value not in (None, "") for c in r)]
    return cab, linhas, 0, idx


def _aaaammdd(v) -> str:
    if hasattr(v, "strftime"):
        return v.strftime("%Y%m%d")
    s = str(v or "").strip()
    if len(s) == 10 and s[2] == "/" and s[5] == "/":
        return s[6:] + s[3:5] + s[:2]
    return "".join(ch for ch in s if ch.isdigit())[:8]


def prazos(inc: dict) -> dict:
    """{cpf: fim do prazo de 30 dias do evento} das vidas da inclusão (17/10/2026: a fila anda pelo prazo)."""
    if not inc.get("xlsx_completo"):
        return {}
    try:
        _, linhas, _, idx = _linhas_xlsx(bytes(inc["xlsx_completo"]))
    except Exception:
        return {}
    pega = lambda l, c: l[idx[c]] if c in idx and idx[c] < len(l) else None  # noqa: E731
    pessoas = []
    for l in linhas:
        cpf = "".join(ch for ch in str(pega(l, "108") or "") if ch.isdigit()).zfill(11)
        tit = "".join(ch for ch in str(pega(l, "132c") or "") if ch.isdigit()).zfill(11)
        rel = str(pega(l, "120") or "").strip().zfill(2)
        pessoas.append({"108": cpf, "132c": tit, "120": rel, "106": _aaaammdd(pega(l, "106")),
                        "132d": _aaaammdd(pega(l, "132d")), "132e": _aaaammdd(pega(l, "132e"))})
    tits = {p["108"]: p for p in pessoas if p["120"] == "01"}
    return {p["108"]: d for p in pessoas if (d := prazo_evento(p, tits.get(p["132c"])))}


def _texto_prazo(inc: dict, pendentes: set) -> str:
    pz = sorted(d for c, d in prazos(inc).items() if c in pendentes)
    if not pz:
        return ""
    dias = (pz[0] - datetime.now().date()).days
    return f"Prazo mais curto: {pz[0]:%d/%m} (" + (f"faltam {dias} dia(s)" if dias >= 0 else f"vencido há {-dias} dia(s)") + ")."


def _xlsx_de(cab: list, linhas: list[list]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(cab)
    for n, l in enumerate(linhas, 1):
        ws.append(l)
        ws.cell(n + 1, 1).value = f"{n:07d}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _so_cpfs(xlsx: bytes, cpfs: set[str]) -> tuple[list, list[list], dict]:
    cab, linhas, _, idx = _linhas_xlsx(xlsx)
    i108 = idx.get("108")
    return cab, [l for l in linhas if str(l[i108] or "").zfill(11) in cpfs], idx


# ── passo 13: quem ficou ativo + gravar no banco ────────────────────────────────────────────────────────
def conferir_ativas(inc: dict, simular: bool) -> tuple[set[str], dict[str, set[str]]]:
    """CPFs que ficaram ativos agora (consulta ao vivo, só dos que faltam) e, por movimentação da inclusão, os que a
    Hapvida CANCELOU depois (19/10/2026, Juracy WP64K: pedido 205/206 cancelado sem documento ficava esquecido)."""
    vidas = {v["cpf"] for v in inc["vidas"] if v.get("cpf")}
    faltam = sorted(vidas - set(inc["ativas"]))
    novas, cancelados = set(), {}
    if not faltam:
        return novas, cancelados
    with coleta_portal.sessao(inc["contrato"]) as (page, headers):
        for cpf in faltam:
            det = coleta_portal.certificados_do_cpf(page, headers, inc["contrato"], cpf)
            if det and any(normalizar((c.get("statusDetail") or {}).get("description") or c.get("_status_lista") or "") == "ATIVO"
                           for c in det):
                novas.add(cpf)
        try:
            cancelados = _cancelados_novos(page, headers, inc, set(faltam) - novas)
        except Exception as e:   # não atrapalha a conferência das ativas
            print(f"[inclusao {inc['id']}] busca de pedido cancelado falhou ({type(e).__name__}: {e})")
    return novas, cancelados


def _movs_da_inclusao(page, headers, inc: dict) -> list[dict]:
    """Movimentações de inclusão do contrato desde a inclusão (o envio e os reenvios do motor)."""
    desde = (_naive(inc["criada_em"]) - timedelta(minutes=10)).isoformat(timespec="seconds")
    return [m for m in coleta_portal.listar_movimentacoes(page, headers, inc["contrato"], 100)
            if (m.get("createdAt") or "") >= desde and not (m.get("uploadName") or "").upper().startswith("EXCLUSAO_")]


def _marca_cancelados(contrato: str, controle: str) -> Path:
    return DADOS / contrato / controle / "cancelados_reanalisados.json"


def _cancelados_novos(page, headers, inc: dict, pendentes: set[str]) -> dict[str, set[str]]:
    """{controle: cpfs} cancelados pela Hapvida que o motor ainda não reanalisou."""
    out = {}
    for m in _movs_da_inclusao(page, headers, inc):
        if m.get("status") not in STATUS_FINAIS:
            continue
        cc = str(m["controlCode"])
        coleta_portal.coletar_movimentacao(page, headers, inc["contrato"], m, DADOS)
        vidas = json.loads((DADOS / inc["contrato"] / cc / "vidas.json").read_text())
        canc = {str((v.get("life") or {}).get("document") or "").zfill(11) for v in vidas if v.get("status") == "Canceled"}
        marca = _marca_cancelados(inc["contrato"], cc)
        ja = set(json.loads(marca.read_text())) if marca.exists() else set()
        novos = (canc & pendentes) - ja
        if novos:
            out[cc] = novos
    return out


def reanalisar_cancelados(inc: dict, cancelados: dict[str, set[str]], simular: bool) -> None:
    """Motor de novo na movimentação: pedido cancelado sem documento sobe de novo pra pendenciar (motor 0.4.6), sem
    aprovação (VIGIA_INCLUSAO); quem já está ativo ou em pedido mais novo sai na conferência do reenvio."""
    import ciclo_criticas
    for cc, cpfs in cancelados.items():
        print(f"[inclusao {inc['id']}] {len(cpfs)} pedido(s) cancelado(s) pela Hapvida na {cc}: motor de novo")
        if simular:
            continue
        try:
            ciclo_criticas.rodar(inc["contrato"], [cc], {"origem": "VIGIA_INCLUSAO", "solicitante": inc["solicitante_upn"],
                                                         "rodada": 1}, notificar=False)
        except Exception as e:
            print(f"[inclusao {inc['id']}] motor na {cc} falhou ({type(e).__name__}: {e}); tento na próxima conferência")
            continue
        marca = _marca_cancelados(inc["contrato"], cc)
        ja = set(json.loads(marca.read_text())) if marca.exists() else set()
        marca.write_text(json.dumps(sorted(ja | cpfs)))
        os.chmod(marca, 0o600)


def gravar_ativas(inc: dict, cpfs: set[str], simular: bool) -> None:
    if not cpfs or not inc.get("xlsx_completo"):
        return
    cab, linhas, _ = _so_cpfs(bytes(inc["xlsx_completo"]), cpfs)
    if not linhas or simular:
        return
    r = requests.post(f"{API}/inclusao/gravar-ativas", json={
        "xlsx_base64": base64.b64encode(_xlsx_de(cab, linhas)).decode("ascii"),
        "nome_arquivo": f"ATIVAS_{inc['contrato']}_{inc['id']}.xlsx"}, timeout=None)
    print(f"[inclusao {inc['id']}] gravadas no banco: {len(linhas)} vida(s) -> {r.status_code} {r.text[:200]}")


# ── passo 17: o motor esgotou? ──────────────────────────────────────────────────────────────────────────
def em_andamento(cur, inc: dict) -> bool:
    """Alguma movimentação do contrato (o envio, reenvio do motor ou correção devolvida) desde a inclusão ainda não
    terminou no vigia #154 (processamento + ciclo de críticas)."""
    cur.execute("""SELECT count(*) AS n FROM movimentacoes_hapvida_acompanhamento
                    WHERE codigo_contrato = %s AND enviado_em >= %s - interval '10 minutes' AND status NOT IN %s""",
                (inc["contrato"], inc["criada_em"], FINAIS_ACOMPANHAMENTO))   # o envio é registrado logo antes da inclusão
    return cur.fetchone()["n"] > 0


def arquivo_pendencias(inc: dict) -> Path | None:
    """O arquivo mais recente que o motor mandou ao pós-vendas pra esse contrato desde a inclusão."""
    desde = _naive(inc["criada_em"]).timestamp()
    achados = [p for p in (SAIDAS / "ciclo").glob(f"*/CORRECAO_MOTOR_{inc['contrato']}_*.xlsx") if p.stat().st_mtime >= desde]
    return max(achados, key=lambda p: p.stat().st_mtime) if achados else None


# ── cartões ──────────────────────────────────────────────────────────────────────────────────────────
def _txt(texto: str, **k) -> dict:
    return {"type": "TextBlock", "text": texto, "wrap": True, **k}


def _rotulo(inc: dict) -> str:
    return f"lote/protocolo {inc['protocolo'] or inc['id']} ({inc['contrato']})"


def cartao_pergunta(inc: dict, faltam: int, prazo_txt: str = "") -> dict:
    return {"type": "AdaptiveCard", "version": "1.4", "body": [
        _txt(f"**Faltam {faltam} vida(s) do {_rotulo(inc)}**", size="Medium")]
        + ([_txt(prazo_txt)] if prazo_txt else []) + [
        _txt("Quer continuar incluindo neste arquivo ou finalizar e mandar essas vidas em outro arquivo depois?")],
        "actions": [
            {"type": "Action.Submit", "title": "Continuar incluindo", "data": {"acao": "inclusao.continuar", "inclusao_id": inc["id"]}},
            {"type": "Action.Submit", "title": "Finalizar", "data": {"acao": "inclusao.finalizar", "inclusao_id": inc["id"]}},
            {"type": "Action.Submit", "title": "Quero o arquivo", "data": {"acao": "inclusao.arquivo", "inclusao_id": inc["id"]}}]}


def cartao_lembrete(inc: dict, faltam: int, prazo_txt: str = "") -> dict:
    return {"type": "AdaptiveCard", "version": "1.4", "body": [
        _txt(f"⏳ **{_rotulo(inc)}**: ainda faltam {faltam} vida(s) esperando correção." + (f" {prazo_txt}" if prazo_txt else "")),
        _txt("Quer o arquivo com o que corrigir?")],
        "actions": [
            {"type": "Action.Submit", "title": "Quero o arquivo", "data": {"acao": "inclusao.arquivo", "inclusao_id": inc["id"]}},
            {"type": "Action.Submit", "title": "Finalizar", "data": {"acao": "inclusao.finalizar", "inclusao_id": inc["id"]}}]}


def cartao_final(inc: dict, ativas: int, encerradas: int, odonto: bool) -> dict:
    corpo = [_txt(f"✅ **Inclusão concluída — {_rotulo(inc)}**", size="Medium", color="Good"),
             {"type": "FactSet", "facts": [{"title": "Vidas ativas", "value": f"{ativas} de {len(inc['vidas'])}"}]
              + ([{"title": "Encerradas sem inclusão", "value": str(encerradas)}] if encerradas else [])}]
    if inc["casamentos"]:
        corpo.append(_txt("**Data de casamento trocada para a inclusão** (atualize com a data correta):", spacing="Medium"))
        corpo += [_txt(f"• {c['nome']} — {c['antiga'] or 'em branco'} → {c['nova']}", spacing="None") for c in inc["casamentos"][:5]]
        if len(inc["casamentos"]) > 5:   # 17/10/2026: cartão curto; a lista completa está no arquivo da inclusão
            corpo.append(_txt(f"… e mais {len(inc['casamentos']) - 5}.", spacing="None"))
    if odonto:
        corpo.append(_txt("🦷 **Odonto Premium**: [baixar o TXT das vidas ativas]({{LINK}})", spacing="Medium"))
    return {"type": "AdaptiveCard", "version": "1.4", "body": corpo}


# ── passo 19: Odonto Premium ─────────────────────────────────────────────────────────────────────────
def codigo_premium(cur, contrato: str) -> str | None:
    cur.execute("""SELECT p.codigo_externo FROM empresa_operadora_codigos s
                     JOIN empresa_operadora_codigos p ON p.empresa_id = s.empresa_id AND p.operadora_id = s.operadora_id
                    WHERE s.codigo_externo = %s AND p.modalidade = 'ODONTO_PREMIUM' ORDER BY p.id LIMIT 1""", (contrato,))
    row = cur.fetchone()
    return row["codigo_externo"] if row else None


def txt_odonto(inc: dict, premium: str) -> tuple[bytes, str] | None:
    """TXT da Hapvida das vidas ativas com 132a = código Odonto Premium e 117 = 9732 (sem gravar no banco)."""
    cpfs = set(inc["ativas"])
    if not cpfs or not inc.get("xlsx_completo"):
        return None
    cab, linhas, idx = _so_cpfs(bytes(inc["xlsx_completo"]), cpfs)
    for l in linhas:
        l[idx["132a"]], l[idx["117"]] = premium, PLANO_ODONTO
    r = requests.post(f"{API}/converter?gravar=false", files={"file": (f"ODONTO_{premium}.xlsx", _xlsx_de(cab, linhas))},
                      timeout=None)
    r.raise_for_status()
    j = r.json()
    return base64.b64decode(j["arquivo_txt_base64"]), f"ODONTO_PREMIUM_{premium}_{inc['contrato']}_{inc['id']}.txt"


def relatorio_final(inc: dict) -> Path | None:
    """O mesmo relatório do /fechar-contrato (49 colunas + Status + Problema), só com as vidas desta inclusão
    (André, 19/10/2026). Falha aqui não impede o card final."""
    try:
        from saneamento.ferramentas import fechar_contrato as fc
        with coleta_portal.sessao(inc["contrato"]) as (page, headers):
            movs = _movs_da_inclusao(page, headers, inc)
            for m in movs:
                coleta_portal.coletar_movimentacao(page, headers, inc["contrato"], m, DADOS)
        meta = [{k: m.get(k) for k in ("controlCode", "uploadName", "status", "createdAt")} for m in movs]
        cpfs = {v["cpf"] for v in inc["vidas"] if v.get("cpf")}
        ativos = {v["cpf"]: v.get("nome") or "" for v in inc["vidas"] if v.get("cpf") in inc["ativas"]}
        sufixo = f"{inc['contrato']}_INCLUSAO_{inc['id']}_{datetime.now():%Y%m%d}"
        rel, st = fc.relatorio(inc["contrato"], meta, ativos, fc.ENTREGA / sufixo, sufixo, so_cpfs=cpfs, odonto=False)
        print(f"[inclusao {inc['id']}] relatório {rel.name}: {st}")
        return rel
    except Exception as e:
        print(f"[inclusao {inc['id']}] relatório final falhou ({type(e).__name__}: {e})")
        return None


def finalizar(cur, inc: dict, motivo: str, simular: bool, avisar_gestor: bool = False) -> None:
    ativas = len(inc["ativas"])
    encerradas = len({v["cpf"] for v in inc["vidas"] if v.get("cpf")}) - ativas
    premium = codigo_premium(cur, inc["contrato"])
    odonto = txt_odonto(inc, premium) if premium and ativas else None
    card = cartao_final(inc, ativas, encerradas, bool(odonto))
    print(f"[inclusao {inc['id']}] FINALIZAR ({motivo}): ativas {ativas}, encerradas {encerradas}, odonto {bool(odonto)}")
    if simular:
        return
    cur.execute("""UPDATE saneamento.inclusao SET status = 'FINALIZADA', finalizada_em = now(), motivo_final = %s
                    WHERE id = %s""", (motivo, inc["id"]))
    corpo = {"texto": f"Inclusão concluída — {_rotulo(inc)}", "cartao": card}
    if odonto:
        corpo.update(arquivo_base64=base64.b64encode(odonto[0]).decode("ascii"), nome_arquivo=odonto[1])
    destinos = [inc["solicitante_upn"]] + ([_gestor(cur)] if avisar_gestor else [])
    for upn in filter(None, dict.fromkeys(destinos)):
        _cora("/notificar-cartao", {"upn": upn, **corpo})
    # o relatório faz login no portal (1–2 min): fora da requisição do botão Finalizar; thread não-daemon (o cron espera)
    threading.Thread(target=enviar_relatorio, args=(dict(inc), list(filter(None, dict.fromkeys(destinos)))),
                     daemon=False).start()


def enviar_relatorio(inc: dict, destinos: list[str]) -> None:
    rel = relatorio_final(inc)
    if not rel:
        return
    for upn in destinos:
        _cora("/notificar-arquivo", {"upn": upn, "texto": f"📋 Relatório da inclusão — {_rotulo(inc)}: uma linha por vida, "
                                     "com Status (Ativo / Em processamento / Não ativo) e o Problema do que não entrou.",
                                     "arquivo_base64": base64.b64encode(rel.read_bytes()).decode("ascii"),
                                     "nome_arquivo": rel.name})


# ── laço principal ─────────────────────────────────────────────────────────────────────────────────────
def processar(cur, inc: dict, agora: datetime, simular: bool) -> None:
    criada = _naive(inc["criada_em"])
    ult = _naive(inc["ultima_conferencia_em"])
    if agora - criada >= PRIMEIRA_CONFERENCIA and (ult is None or agora - ult >= CONFERIR_A_CADA):
        try:
            novas, cancelados = conferir_ativas(inc, simular)
        except Exception as e:
            print(f"[inclusao {inc['id']}] conferência falhou ({type(e).__name__}: {e}); tento na próxima rodada")
            novas, cancelados = None, {}
        if cancelados:
            reanalisar_cancelados(inc, cancelados, simular)
        if novas is not None:
            gravar_ativas(inc, novas, simular)
            agora_iso = agora.isoformat(timespec="seconds")
            inc["ativas"] = {**inc["ativas"], **{c: agora_iso for c in novas}}
            if not simular:
                cur.execute("UPDATE saneamento.inclusao SET ativas = %s, ultima_conferencia_em = now() WHERE id = %s",
                            (json.dumps(inc["ativas"]), inc["id"]))
    pendentes = {v["cpf"] for v in inc["vidas"] if v.get("cpf")} - set(inc["ativas"])
    faltam = len(pendentes)
    if faltam == 0:
        return finalizar(cur, inc, "TODAS_ATIVAS", simular)
    if not no_expediente(agora) or em_andamento(cur, inc):
        return
    st = inc["status"]
    if st == "ABERTA":
        print(f"[inclusao {inc['id']}] pergunta Continuar/Finalizar (faltam {faltam})")
        if not simular and _cora("/notificar-cartao", {"upn": inc["solicitante_upn"], "texto": f"Faltam {faltam} vida(s) do {_rotulo(inc)}",
                                                     "cartao": cartao_pergunta(inc, faltam, _texto_prazo(inc, pendentes))}):
            cur.execute("""UPDATE saneamento.inclusao SET status = 'PERGUNTADA', perguntada_em = now(),
                           ultimo_lembrete_em = now() WHERE id = %s""", (inc["id"],))
        return
    base = _naive(inc["decisao_em"] if st == "CONTINUAR" and inc["decisao_em"] else inc["perguntada_em"]) or criada
    # 16/10/2026 (André): os prazos de 4 e 7 dias são em dias ÚTEIS — dia útil = um expediente inteiro (seg–sex)
    dias = horas_uteis(base, agora) / (EXPEDIENTE[1] - EXPEDIENTE[0])
    # 17/10/2026 (André): com data de evento no arquivo, o relógio é o prazo de 30 dias do evento — não encerra antes
    # dele vencer e avisa o gestor 10 dias antes; sem data, vale o de sempre (4 e 7 dias úteis sem resposta)
    pz = sorted(d for c, d in prazos(inc).items() if c in pendentes)
    hoje = agora.date()
    vencido = (hoje > pz[-1]) if pz else dias >= DIAS_FINALIZA
    perto = ((pz[0] - hoje).days <= AVISO_GESTOR_ANTES_DO_PRAZO) if pz else dias >= DIAS_GESTOR
    if vencido and dias >= DIAS_FINALIZA:
        return finalizar(cur, inc, "PRAZO", simular, avisar_gestor=True)
    if perto and not inc["gestor_avisado_em"]:
        gestor = _gestor(cur)
        print(f"[inclusao {inc['id']}] perto do prazo/sem resposta: pergunta de novo com cópia ao gestor ({bool(gestor)})")
        if not simular:
            for upn in filter(None, dict.fromkeys([inc["solicitante_upn"], gestor])):
                _cora("/notificar-cartao", {"upn": upn, "texto": f"Pendente — {_rotulo(inc)}",
                                            "cartao": cartao_pergunta(inc, faltam, _texto_prazo(inc, pendentes))})
            cur.execute("UPDATE saneamento.inclusao SET gestor_avisado_em = now(), ultimo_lembrete_em = now() WHERE id = %s",
                        (inc["id"],))
        return
    if st == "CONTINUAR" and horas_uteis(_naive(inc["ultimo_lembrete_em"]) or base, agora) >= LEMBRETE_HORAS_UTEIS:
        print(f"[inclusao {inc['id']}] lembrete (faltam {faltam})")
        if not simular and _cora("/notificar-cartao", {"upn": inc["solicitante_upn"], "texto": f"Lembrete — {_rotulo(inc)}",
                                                     "cartao": cartao_lembrete(inc, faltam, _texto_prazo(inc, pendentes))}):
            cur.execute("UPDATE saneamento.inclusao SET ultimo_lembrete_em = now() WHERE id = %s", (inc["id"],))


def _abertas(cur, inclusao_id: int | None = None) -> list[dict]:
    cur.execute("""SELECT * FROM saneamento.inclusao WHERE status <> 'FINALIZADA' AND (%s::bigint IS NULL OR id = %s)
                   ORDER BY id""", (inclusao_id, inclusao_id))
    return [dict(r) for r in cur.fetchall()]


def run(simular: bool = False) -> None:
    if not simular:
        try:
            import ciclo_criticas
            ciclo_criticas.retentar_reenvios()
        except Exception as e:
            print(f"[reenvio] retentativa falhou: {type(e).__name__}: {e}")
    conn = _conn()
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            for inc in _abertas(cur):
                try:
                    processar(cur, inc, datetime.now(), simular)
                    conn.commit()
                except Exception as e:
                    conn.rollback()
                    print(f"[inclusao {inc['id']}] erro: {type(e).__name__}: {e}")
    finally:
        conn.close()


# ── decisão do cartão (app.py /inclusao/decisao) ─────────────────────────────────────────────────────────
def decidir(inclusao_id: int, acao: str, por: str) -> dict:
    conn = _conn()
    try:
        with conn, conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            abertas = _abertas(cur, inclusao_id)
            if not abertas:
                raise Recusa("essa inclusão já foi finalizada — nada a fazer")
            inc = abertas[0]
            autorizados = {u.lower() for u in (inc["solicitante_upn"], _gestor(cur), ADMIN) if u}
            if por.lower() not in autorizados:
                raise Recusa("só quem fez a inclusão (ou o gestor do pós-vendas) decide esta inclusão")
            if acao == "arquivo":
                arq = arquivo_pendencias(inc)
                if not arq:
                    return {"texto": f"Não achei arquivo de correção para o {_rotulo(inc)} — o motor ainda não mandou pendência."}
                return {"texto": f"📋 Arquivo do {_rotulo(inc)}: corrija só as células em **vermelho** e me devolva "
                                 "escrevendo **correção motor de críticas**.",
                        "arquivo_base64": base64.b64encode(arq.read_bytes()).decode("ascii"), "nome_arquivo": arq.name}
            if acao == "continuar":
                cur.execute("""UPDATE saneamento.inclusao SET status = 'CONTINUAR', decisao_em = now(), decidida_por = %s,
                               ultimo_lembrete_em = now(), gestor_avisado_em = NULL WHERE id = %s""", (por, inc["id"]))
                return {"texto": f"Ok, sigo esperando a correção do {_rotulo(inc)}. Te lembro a cada 8h úteis."}
            if acao == "finalizar":
                cur.execute("UPDATE saneamento.inclusao SET decisao_em = now(), decidida_por = %s WHERE id = %s", (por, inc["id"]))
                finalizar(cur, inc, "BOTAO", simular=False)
                return {"texto": f"Finalizado: as vidas que faltavam ficam como não incluídas neste arquivo. O resumo vai no card."}
            raise Recusa("ação desconhecida")
    finally:
        conn.close()


if __name__ == "__main__":
    run(simular="--simular" in sys.argv)
