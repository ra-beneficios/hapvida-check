"""
Procedimento mensal de faturamento Hapvida (Andre, 28/09/2026): baixa os
documentos de faturamento do Grupo Saude pra pasta compartilhada e, no fim,
audita TODAS as pastas do mes (Saude + Legado/Cowork) dizendo o que falta.

Cron no host (00:01):
- dia 26 00:01 (--modo=completo): baixa tudo o que ja estiver disponivel no portal.
- ultimo dia do mes e dia 04, 00:01 (--modo=faltantes): so' o que ainda falta.

Competencia = mes do vencimento: rodando do dia 20 em diante, e' o mes
seguinte (28/09 -> OUT26); antes do dia 20, o proprio mes (04/10 -> OUT26).

Legado (portal antigo), 07/10/2026: contrato (Odonto ou Saude) com credencial no
Bitwarden (segredo hapvida-legado, copiado pelo sync_bitwarden_legado.py no
host) e' baixado aqui pelo legado_auto.py; o resto continua com o Cowork
(handoff). Lista viva em _status/LEGADO_ASSUMIDO_CODE.json.

Nada aqui manda e-mail pra cliente - so' salva arquivo (enviar_email=False)
e gera o relatorio de auditoria em faturas/_status/ + handoff pro Cowork
com o que falta do lado Legado.

Lista de ativos e analitico de copart PDF sao gerados em fila pela Hapvida:
pedidos numa execucao, coletados na seguinte (26 -> ultimo dia -> 04).

Pausar os downloads: criar _status/FATURAMENTO_PAUSADO (so' roda a auditoria).

Uso: python3 faturamento_mensal.py --modo=completo|faltantes [--competencia=OUT26]
"""
import csv
import io
import json
import os
import re
import sys
import time
import unicodedata
from datetime import date, datetime, timedelta

import httpx
import openpyxl
import openpyxl.styles
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
import psycopg2
from playwright.sync_api import sync_playwright

import boletos_saude
import censo_saude
import enviar_documentos
import legado_auto
import padrao_faturamento as pf
import scraper
import secrets_store
from censo_saude import login_saude

BASE = enviar_documentos.SHAREPOINT_FATURAS_BASE
# Mesmo mount do /sharepoint/faturas, dois niveis acima - onde mora o
# protocolo _handoff/ do Cowork. Montado a parte no docker-compose.yml.
HANDOFF_ENTRADA = "/projeto_ra/_handoff/entrada"
PLANILHA_LEGADO = "/projeto_ra/operadoras/Access/Hapvida/empresas_hapvida.xlsx"

# DAVJV/B6PJX so' abrem com a conta do Marcos/Vinicius (#147) - mesma regra
# do endpoint /enviar-documentos-contrato.
CODIGOS_CONTA_MARCO = ("DAVJV", "B6PJX")
# Contratos ativos SEM faturamento proprio pra baixar/enviar (Andre, 02/10/2026):
# - 9WLZU/9GBKA (Boaventura): fora do envio por decisao gerencial (so' vidas, #150)
# - ENUHC (LH) e DLIHM (Heftos): ativos, mas sem movimentacao e sem fatura
# - 340ZM (Sintramon): ativo, fatura paga dentro de outra nota
# Continuam ativos no banco - so' saem do escopo de faturamento.
SEM_FATURAMENTO = {"9WLZU", "9GBKA", "ENUHC", "DLIHM", "340ZM"}
TENTATIVAS_DOWNLOAD = 3
PAUSA_ENTRE_TENTATIVAS = 60
censo_saude.POLL_TIMEOUT_SECONDS = 420  # 0KFXW (~50 coligadas) estoura 240s

MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]
COLUNAS_CENTAVOS = {"mensalidade", "adicional", "taxa_adesao", "desconto", "cobrado"}


def log(msg: str) -> None:
    print(f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {msg}", flush=True)


def competencia_padrao(hoje: date) -> str:
    alvo = (hoje.replace(day=1) + timedelta(days=32)) if hoje.day >= 20 else hoje
    return f"{MESES[alvo.month - 1]}{alvo.year % 100:02d}"


def competencia_para_mes(comp: str) -> tuple[int, int]:
    return MESES.index(comp[:3]) + 1, 2000 + int(comp[3:])


def competencia_anterior(comp: str) -> str:
    mes, ano = competencia_para_mes(comp)
    mes, ano = (12, ano - 1) if mes == 1 else (mes - 1, ano)
    return f"{MESES[mes - 1]}{ano % 100:02d}"


def _normaliza(s: str) -> str:
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().upper()


# ---------------------------------------------------------------- pastas

def pasta_contrato(codigo: str, comp: str, empresa_nome: str | None = None) -> str:
    """Reaproveita a pasta de empresa que ja existe pro codigo (o nome varia
    entre Cowork e banco, ex: 'C-MMIND CONST MANUT' x 'C-MMIND CONST, MANUT'
    - criar outra geraria pasta duplicada). Se houver mais de uma, usa a que
    ja tem a competencia; senao a modificada mais recentemente."""
    candidatas = [os.path.join(BASE, e, codigo) for e in os.listdir(BASE)
                  if os.path.isdir(os.path.join(BASE, e, codigo))]
    if candidatas:
        com_comp = [c for c in candidatas if os.path.isdir(os.path.join(c, comp))]
        escolhida = max(com_comp or candidatas, key=os.path.getmtime)
        return os.path.join(escolhida, comp)
    nome = empresa_nome or enviar_documentos._empresa_nome(codigo)
    return os.path.join(BASE, enviar_documentos._nome_pasta_empresa(nome), codigo, comp)


def arquivos(pasta: str) -> list[str]:
    return sorted(os.listdir(pasta)) if os.path.isdir(pasta) else []


_RE_NUMERO_NFSE = re.compile(r"N[ÚU]MERO DA NFS-?E\D{0,200}?(\d{6,12})", re.I | re.S)


_RE_COMPETENCIA_NF = re.compile(r"COMPETENCIA\s+(\d{2})(\d{4})")


def nfs_do_mes(pasta: str, nomes: list[str], comp: str) -> tuple[list[str], list[str]]:
    """Separa as NFs da competencia das de outro mes, pela "COMPETENCIA MMAAAA"
    que a DANFSe do Legado traz no texto (Z8TUR 01/10/2026: NF de SET26 salva
    como NFE_SAUDE_Z8TUR_OUT26). Sem competencia legivel, conta como do mes."""
    import pypdf
    mes, ano = competencia_para_mes(comp)
    do_mes, outro = [], []
    for n in nomes:
        try:
            texto = "".join(p.extract_text() or "" for p in pypdf.PdfReader(os.path.join(pasta, n)).pages)
            m = _RE_COMPETENCIA_NF.search(texto)
        except Exception:
            m = None
        (outro if m and (int(m.group(1)), int(m.group(2))) != (mes, ano) else do_mes).append(n)
    return do_mes, outro


def conteudos_distintos(pasta: str, nomes: list[str]) -> set[str]:
    """Identidade de cada documento: o numero da NFS-e lido do PDF (DANFSe do
    Legado traz "NUMERO DA NFS-E"); sem numero, o hash do arquivo. A mesma
    nota salva com dois nomes conta uma vez (Z8TUR, 01/10/2026 - os PDFs
    tinham bytes diferentes)."""
    import hashlib
    import pypdf
    ids = set()
    for n in nomes:
        caminho = os.path.join(pasta, n)
        numero = None
        try:
            texto = "".join(p.extract_text() or "" for p in pypdf.PdfReader(caminho).pages)
            m = _RE_NUMERO_NFSE.search(texto)
            numero = m and m.group(1)
        except Exception:
            pass
        if not numero:
            with open(caminho, "rb") as f:
                numero = "md5:" + hashlib.md5(f.read()).hexdigest()
        ids.add(numero)
    return ids


def tem(nomes: list[str], prefixo: str, ext: str | None = None) -> bool:
    return any(_normaliza(n).startswith(prefixo) and (ext is None or n.lower().endswith(ext)) for n in nomes)


# ---------------------------------------------------------------- escopo

def contratos_saude(comp: str) -> dict[str, set[str]]:
    """codigo -> tipos de fatura (SAUDE/COPARTICIPACAO) com vencimento na
    competencia. So' entra quem tem MENSALIDADE no mes - fatura so' de
    copart (ex: ENUHC, sem mensalidade desde AGO26) fica de fora, senao o
    envio_documentos gravaria na pasta do mes da ultima mensalidade."""
    mes, ano = competencia_para_mes(comp)
    conn = psycopg2.connect(dbname="lifetools")
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT eoc.codigo_externo, f.bill_type
                FROM faturas f JOIN empresa_operadora_codigos eoc ON eoc.id = f.empresa_operadora_codigo_id
                WHERE eoc.operadora_id = 1
                  AND extract(month FROM f.vencimento) = %s AND extract(year FROM f.vencimento) = %s
                """,
                (mes, ano),
            )
            por_codigo: dict[str, set[str]] = {}
            for codigo, bill_type in cur.fetchall():
                tipo = "SAUDE" if bill_type == "MENSALIDADE PLANO SAUDE" else "COPARTICIPACAO"
                por_codigo.setdefault(codigo, set()).add(tipo)
    finally:
        conn.close()
    return {c: t for c, t in por_codigo.items() if "SAUDE" in t and c not in SEM_FATURAMENTO}


def contratos_legado() -> dict[str, str]:
    """codigo -> empresa, aba 'Contratos (55)' da planilha do Cowork, so'
    Status=Ativo (os 7 Desativados do AVISO_CLAUDE_CODE.md ficam fora).
    Le so' Contrato/Empresa/Status - nunca as colunas de usuario/senha."""
    wb = openpyxl.load_workbook(PLANILHA_LEGADO, read_only=True)
    ws = next(ws for ws in wb if ws.title.startswith("Contratos"))
    linhas = ws.iter_rows(values_only=True)
    cab = [str(h or "").strip() for h in next(linhas)]
    i_cod, i_emp, i_st = cab.index("Contrato"), cab.index("Empresa (Razão Social)"), cab.index("Status")
    ativos = {}
    for r in linhas:
        if r[i_cod] and str(r[i_st] or "").strip().lower() == "ativo":
            if str(r[i_cod]).strip() not in SEM_FATURAMENTO:
                ativos[str(r[i_cod]).strip()] = str(r[i_emp] or "").strip()
    return ativos


def status_legado(comp: str) -> dict[str, dict]:
    """Do CSV de status do Cowork: por codigo, se ja existe boleto/NF com
    numero emitido e vencimento na competencia - separa 'faltando na pasta'
    de 'Hapvida ainda nao emitiu'."""
    mes, ano = competencia_para_mes(comp)
    pasta = os.path.join(BASE, "_status")
    arquivos_status = sorted(
        (os.path.join(pasta, n) for n in os.listdir(pasta)
         if n.startswith("STATUS_FATURAS_") and not n.startswith("STATUS_FATURAS_CODE_") and n.endswith(".csv")),
        key=os.path.getmtime,
    ) if os.path.isdir(pasta) else []
    # 07/10/2026: contratos do Legado que o Code assumiu (legado_auto) tem status proprio,
    # que vale no lugar do Cowork pra esses codigos.
    do_code = os.path.join(pasta, f"STATUS_FATURAS_CODE_{comp}.csv")
    fontes = arquivos_status[-1:] + ([do_code] if os.path.exists(do_code) else [])
    res: dict[str, dict] = {}
    for fonte in fontes:
        with open(fonte, encoding="utf-8-sig") as f:
            linhas = list(csv.DictReader(f, delimiter=";"))
        if fonte == do_code:
            for c in {l["Codigo"] for l in linhas}:
                res.pop(c, None)
        _acumular_status(res, linhas, mes, ano)
    return res


def _acumular_status(res: dict, linhas: list[dict], mes: int, ano: int) -> None:
    for linha in linhas:
        try:
            d, m, a = linha["Vencimento"].split("/")
        except (ValueError, KeyError, AttributeError):
            continue
        if int(m) != mes or 2000 + int(a[-2:]) != ano:
            continue
        info = res.setdefault(linha["Codigo"], {"emitido": False, "tipos": set(), "boletos": set(), "valores": set()})
        numero = (linha.get("NumeroDocumento") or "").strip()
        tipo = linha.get("TipoDocumento", "").strip().upper()
        if numero:
            info["emitido"] = True
            info["tipos"].add(tipo)
            if tipo == "BOLETO":
                info["boletos"].add(numero)
                v = pf._float_br(linha.get("Valor") or "")
                if v is not None:
                    info["valores"].add(v)


# ---------------------------------------------------------------- saude

def valores_saude(codigo: str, comp: str) -> dict[str, set[float]]:
    """Valor de cada boleto do contrato na competencia, por plano (SAUDE/COPARTICIPACAO)."""
    mes, ano = competencia_para_mes(comp)
    conn = psycopg2.connect(dbname="lifetools")
    try:
        with conn.cursor() as cur:
            cur.execute(
                """SELECT f.bill_type, f.valor FROM faturas f
                   JOIN empresa_operadora_codigos eoc ON eoc.id = f.empresa_operadora_codigo_id
                   WHERE eoc.codigo_externo = %s
                     AND extract(month FROM f.vencimento) = %s AND extract(year FROM f.vencimento) = %s""",
                (codigo, mes, ano))
            out: dict[str, set[float]] = {}
            for bill_type, valor in cur.fetchall():
                plano = "SAUDE" if bill_type == "MENSALIDADE PLANO SAUDE" else "COPARTICIPACAO"
                out.setdefault(plano, set()).add(float(valor))
            return out
    finally:
        conn.close()


def faltando_saude(codigo: str, tipos: set[str], comp: str, conferir: bool = False) -> list[str]:
    """Padrao NOV26 (padrao_faturamento): arquivos de envio + CSV oficiais (fonte
    do XLSX/PDF e da conferencia). Com conferir=True, inclui total x boleto."""
    pasta = pasta_contrato(codigo, comp)
    nomes = arquivos(pasta)
    copart = "COPARTICIPACAO" in tipos
    falta = pf.faltando(nomes, codigo, comp, "SAUDE", copart)
    falta += [n for n in pf.faltando_fontes(nomes, codigo, comp, "SAUDE", copart) if n not in falta]
    if conferir:
        falta += pf.conferir_valores(pasta, codigo, comp, "SAUDE", copart, valores_saude(codigo, comp))
    return falta


def _so_documentos_portal(falta: list[str]) -> bool:
    """Falta algo que so' o portal da' (boleto, NF, relatorio PDF/CSV)? XLSX e
    PDF da lista sao gerados aqui; a lista em si vem do salvar_listas_ativos."""
    return any(f.startswith(("BOLETO_", "NFE_", "RELATORIO_ANALITICO_")) and not f.lower().endswith(".xlsx")
               for f in falta)


def _lista_csv(codigo: str, comp: str) -> str:
    return pf.nome("LISTA_ATIVOS", "SAUDE", codigo, comp, "csv")


def tem_lista(codigo: str, comp: str) -> bool:
    return _lista_csv(codigo, comp).upper() in {n.upper() for n in arquivos(pasta_contrato(codigo, comp))}


def baixar_documentos_saude(codigos: dict[str, set[str]], comp: str, modo: str, creds: dict) -> dict[str, list[str]]:
    avisos_por_codigo: dict[str, list[str]] = {}
    pendentes = list(codigos) if modo == "completo" else [
        c for c in codigos if _so_documentos_portal(faltando_saude(c, codigos[c], comp))
    ]
    for tentativa in range(1, TENTATIVAS_DOWNLOAD + 1):
        if not pendentes:
            break
        if tentativa > 1:
            log(f"tentativa {tentativa}: {len(pendentes)} contrato(s) ainda incompletos, aguardando {PAUSA_ENTRE_TENTATIVAS}s")
            time.sleep(PAUSA_ENTRE_TENTATIVAS)
        for codigo in pendentes:
            conta = creds["marco"] if codigo in CODIGOS_CONTA_MARCO else creds["carol"]
            try:
                r = enviar_documentos.montar_e_enviar(
                    codigo, "", os.environ["GCP_PROJECT"], conta["usuario"], conta["senha"], enviar_email=False
                )
                avisos_por_codigo[codigo] = r.get("avisos") or ([r["detalhe"]] if r.get("detalhe") else [])
                log(f"{codigo}: {len(r.get('anexos', []))} arquivo(s), {len(avisos_por_codigo[codigo])} aviso(s)")
            except Exception as e:
                avisos_por_codigo[codigo] = [f"erro: {e}"]
                log(f"{codigo}: ERRO {e}")
        # montar_e_enviar so' sobrescreve o que conseguiu baixar - o que ja
        # estava na pasta de uma tentativa anterior continua la.
        pendentes = [c for c in pendentes if _so_documentos_portal(faltando_saude(c, codigos[c], comp))]
    return avisos_por_codigo


_RE_EXCEL_TEXTO_FORCADO = re.compile(r'="([^"]*)"')


def _texto_censo(bruto: bytes) -> str:
    """Mesma normalizacao do script 160 (achado 18/09): export as vezes vem
    em latin-1, com linha em branco antes do cabecalho e campo embrulhado
    em ="valor" (Excel forcando texto)."""
    try:
        texto = bruto.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = bruto.decode("latin-1")
    texto = _RE_EXCEL_TEXTO_FORCADO.sub(r"\1", texto)
    return "\n".join(l for l in texto.splitlines() if l.strip())


ESTADO_LISTAS = os.path.join(BASE, "_status", "LISTAS_ATIVOS_ESTADO.json")


def salvar_listas_ativos(codigos: dict[str, set[str]], comp: str, modo: str, creds: dict) -> list[str]:
    """Lista de ativos OFICIAL do portal (relatorio Full de beneficiarios da
    mae, mesmo arquivo do script 160), quebrada por contrato em
    LISTA_ATIVOS_SAUDE_{codigo}_{comp}.csv (fonte do XLSX e do PDF). So' grava arquivo; NAO mexe no banco
    de vidas (isso continua sendo o 160, manual).

    Assincrono desde 01/10/2026: a Hapvida chega a levar mais de 40 min pra
    gerar (0KFXW estoura qualquer espera razoavel) e o status do relatorio e'
    por USUARIO, nao por empresa - um pedido atrasado aparece como resultado
    do pedido seguinte. Entao: (1) se um relatorio pedido numa execucao
    anterior ficou pronto, coleta; (2) pede o das maes que ainda tem filho
    sem lista e espera ate' POLL_TIMEOUT_SECONDS; se nao sair, fica registrado
    em ESTADO_LISTAS e a proxima execucao coleta. O CSV traz o codigo de cada
    linha (coluna EMPRESA), entao o arquivo se identifica sozinho: so' grava
    linha de contrato que ainda falta."""
    alvo = set(codigos) if modo == "completo" else {c for c in codigos if not tem_lista(c, comp)}
    if not alvo:
        return []
    gravados: set[str] = set()

    def faltam() -> set[str]:
        # completo regrava todas; faltantes so' as que nao estao na pasta
        if modo == "completo":
            return alvo - gravados
        return {c for c in alvo if not tem_lista(c, comp)}

    mae_de = _maes_dos_codigos(alvo)
    estado = _ler_estado_listas()
    erros = []
    grupos = [("carol", censo_saude.MAES_SAUDE)]
    if alvo & set(CODIGOS_CONTA_MARCO):
        grupos.append(("marco", ["DAVJV"]))
    for conta, maes in grupos:
        st = estado.setdefault(conta, {})
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            page = browser.new_page()
            try:
                auth = login_saude(page, creds[conta]["usuario"], creds[conta]["senha"])
            finally:
                browser.close()
        coletadas: set[str] = set()  # mae ja' coletada nesta execucao nao e' pedida de novo
        try:
            ct = _status_relatorio_full(auth, esperar=True)
            link = ct.get("downloadLink")
            if ct.get("status") == "Processed" and link and st.get("pedido") and link != st.get("link_consumido"):
                gravados |= _gravar_listas(_baixar(link), alvo, comp)
                log(f"lista ativos: coletado relatório pedido antes (mãe {st['pedido']}, {st.get('pedido_em')})")
                coletadas.add(st["pedido"])
                st.update(pedido=None, link_consumido=link)
                _salvar_estado_listas(estado)
            for mae in maes:
                pendentes = faltam()
                if mae in coletadas or not any(mae_de.get(c) in (mae, None) for c in pendentes):
                    continue
                if ct.get("status") == "Processing":
                    erros.append(f"Lista de ativos: relatório da mãe {st.get('pedido')} ainda em geração na Hapvida — coletado na próxima execução")
                    break
                censo_saude.criar_sessao_empresa(auth, mae)
                censo_saude.gerar_relatorio(auth, mae, "Full")
                st.update(pedido=mae, pedido_em=datetime.now().isoformat(timespec="minutes"))
                _salvar_estado_listas(estado)
                try:
                    link = censo_saude._poll_status(auth, "Full", link)
                except censo_saude.RelatorioTimeoutError:
                    erros.append(f"Lista de ativos (mãe {mae}): pedida, ainda em geração na Hapvida — coletada na próxima execução")
                    break
                gravados |= _gravar_listas(_baixar(link), alvo, comp)
                st.update(pedido=None, link_consumido=link)
                _salvar_estado_listas(estado)
                ct = {"status": "Processed", "downloadLink": link}
        except Exception as e:
            erros.append(f"Lista de ativos (conta {conta}): {e}")
            log(f"lista ativos {conta}: ERRO {e}")
    resto = faltam()
    if resto:
        erros.append(f"Lista de ativos ainda faltando: {', '.join(sorted(resto))}")
    return erros


def _maes_dos_codigos(codigos: set[str]) -> dict[str, str | None]:
    """codigo -> mae (o proprio codigo se for mae; None se nao souber)."""
    conn = psycopg2.connect(dbname="lifetools")
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT codigo_externo, codigo_contrato_mae FROM empresa_operadora_codigos "
                        "WHERE operadora_id = 1 AND codigo_externo = ANY(%s)", (list(codigos),))
            maes = set(censo_saude.MAES_SAUDE) | {"DAVJV"}
            return {c: (m or (c if c in maes else None)) for c, m in cur.fetchall()}
    finally:
        conn.close()


def _status_relatorio_full(auth: dict, esperar: bool) -> dict:
    url = f"{censo_saude.API_BASE}/beneficiaries/status-report?type=Full"
    inicio = time.time()
    while True:
        r = httpx.get(url, headers=censo_saude._headers_http(auth), timeout=30)
        r.raise_for_status()
        ct = r.json().get("content", {}) or {}
        if not esperar or ct.get("status") != "Processing" or time.time() - inicio > censo_saude.POLL_TIMEOUT_SECONDS:
            return ct
        time.sleep(censo_saude.POLL_INTERVAL_SECONDS)


def _baixar(link: str) -> bytes:
    r = httpx.get(link, timeout=120)
    r.raise_for_status()
    return r.content


def _gravar_listas(bruto: bytes, alvo: set[str], comp: str) -> set[str]:
    leitor = csv.reader(io.StringIO(_texto_censo(bruto)), delimiter=";", quotechar='"')
    cab = next(leitor)
    por_codigo: dict[str, list] = {}
    for linha in leitor:
        if linha and linha[0] in alvo:
            por_codigo.setdefault(linha[0], []).append(linha)
    for codigo, linhas in por_codigo.items():
        pasta = pasta_contrato(codigo, comp)
        os.makedirs(pasta, exist_ok=True)
        with open(os.path.join(pasta, _lista_csv(codigo, comp)), "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f, delimiter=";", quotechar='"', quoting=csv.QUOTE_ALL)
            w.writerow(cab)
            w.writerows(linhas)
    log(f"lista ativos: {len(por_codigo)} contrato(s) gravados ({', '.join(sorted(por_codigo)) or '-'})")
    return set(por_codigo)


def _ler_estado_listas() -> dict:
    try:
        with open(ESTADO_LISTAS, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def _salvar_estado_listas(estado: dict) -> None:
    with open(ESTADO_LISTAS, "w", encoding="utf-8") as f:
        json.dump(estado, f, ensure_ascii=False, indent=1)


# ---------------------------------------------------------------- xlsx

def _num(v: str):
    try:
        return int(v)
    except ValueError:
        return v


def gerar_xlsx_relatorios(comp: str) -> int:
    """Derivados de todo CSV oficial da pasta da competencia (Saude e Legado):
    XLSX de relatorio/lista (Andre, 01/10/2026) e PDF da lista de ativos
    (regra NOV26, 05/10/2026). O portal nao oferece XLSX: o analitico so' sai
    em PDF/CSV/TXT (XLSX/XLS da 500, testado 28/09/2026) e a lista de ativos
    so' em CSV/tela."""
    gerados = 0
    pastas = []
    for empresa in os.listdir(BASE):
        dir_emp = os.path.join(BASE, empresa)
        if empresa.startswith("_") or not os.path.isdir(dir_emp):
            continue
        pastas += [os.path.join(dir_emp, codigo, comp) for codigo in os.listdir(dir_emp)]
    for pasta in pastas:
        gerados += gerar_xlsx_pasta(pasta)
    gerados += gerar_pdf_listas(pastas)
    return gerados


def _desatualizado(origem: str, destino: str) -> bool:
    return not os.path.exists(destino) or os.path.getmtime(destino) < os.path.getmtime(origem)


def gerar_xlsx_pasta(pasta: str) -> int:
    gerados = 0
    for nome in arquivos(pasta):
        up = nome.upper()
        if not nome.lower().endswith(".csv") or up.startswith(("BENEFICIARIOS_ATIVOS", "COPARTICIPACAO_ANALITICO")):
            continue
        destino = os.path.join(pasta, nome[:-4] + ".xlsx")
        origem = os.path.join(pasta, nome)
        if not _desatualizado(origem, destino):
            continue
        try:
            if up.startswith("RELATORIO_ANALITICO_COPARTICIPACAO"):
                _csv_copart_para_xlsx(origem, destino)
            elif up.startswith("RELATORIO_ANALITICO"):
                _csv_para_xlsx(origem, destino)
            else:
                _csv_generico_para_xlsx(origem, destino)
            gerados += 1
        except Exception as e:
            log(f"xlsx {nome}: ERRO {e}")
    return gerados


def _csv_copart_para_xlsx(origem: str, destino: str) -> None:
    """Analitico de coparticipacao oficial (CSV da fatura) -> XLSX: cabecalho
    da Hapvida em cima, valores (colunas VL*) como numero em reais, CPF/senha
    como texto, linha de total no fim."""
    with open(origem, "rb") as f:
        linhas = _texto_censo(f.read()).splitlines()
    i_cab = next((i for i, l in enumerate(linhas) if "VL PARTICIPACAO" in l.upper()),
                 next(i for i, l in enumerate(linhas) if l.count(";") >= 5))
    cab = [c.strip() for c in linhas[i_cab].split(";")]
    while cab and not cab[-1]:
        cab.pop()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Coparticipacao"
    for l in linhas[:i_cab]:
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", l.strip())])
        ws.cell(ws.max_row, 1).font = openpyxl.styles.Font(bold=True)
    ws.append([])
    ws.append(cab)
    linha_cab = ws.max_row
    for c in ws[linha_cab]:
        c.font = openpyxl.styles.Font(bold=True)
    valor = {i for i, c in enumerate(cab) if c.upper().startswith("VL")}
    i_part = next((i for i, c in enumerate(cab) if c.upper() == "VL PARTICIPACAO"), None)
    total, n = 0.0, 0
    for valores in csv.reader(io.StringIO("\n".join(linhas[i_cab + 1:])), delimiter=";", quotechar='"'):
        valores = [ILLEGAL_CHARACTERS_RE.sub("", v.strip()) for v in valores[:len(cab)]]
        if not any(valores):
            continue
        linha = []
        for i, v in enumerate(valores):
            if i in valor:
                f = pf._float_br(v)
                linha.append(f if f is not None else v)
            elif cab[i].upper() == "QUANTIDADE" and v.isdigit():
                linha.append(int(v))
            else:
                linha.append(v)
        ws.append(linha)
        n += 1
        for i in valor:
            ws.cell(ws.max_row, i + 1).number_format = "#,##0.00"
        if i_part is not None and isinstance(linha[i_part] if i_part < len(linha) else None, float):
            total += linha[i_part]
    if i_part is not None:
        ws.append([])
        ws.append(["TOTAL"] + [None] * (i_part - 1) + [round(total, 2)])
        ws.cell(ws.max_row, 1).font = openpyxl.styles.Font(bold=True)
        ws.cell(ws.max_row, i_part + 1).number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(linha_cab + 1, 1)
    if n:
        ws.auto_filter.ref = f"A{linha_cab}:{openpyxl.utils.get_column_letter(len(cab))}{linha_cab + n}"
    wb.save(destino)


# Colunas da lista de ativos que vao pro PDF (nome no PDF -> nomes possiveis no CSV
# do portal novo, 45 colunas, ou do Legado, 12 colunas). O XLSX leva todas.
_COLUNAS_LISTA_PDF = [
    ("Código", ("CODIGO USUARIO",)),
    ("Matrícula", ("MATRICULA",)),
    ("Beneficiário", ("USUARIO",)),
    ("CPF", ("CPF_DEPENDENTE", "CPF")),
    ("Tipo", ("TIPO USUARIO", "TIPO_USUARIO")),
    ("Parentesco", ("GRAU PARENTESCO",)),
    ("Plano", ("PLANO",)),
    ("Nascimento", ("DATA_NASCIMENTO",)),
    ("Admissão", ("DT_ADMISSAO",)),
]


def gerar_pdf_listas(pastas: list[str]) -> int:
    """LISTA_ATIVOS_{PLANO}_{COD}_{MES}.pdf a partir do CSV oficial do mes
    (Andre, 05/10/2026: lista de ativos vai em XLSX e PDF). So' gera quando o
    PDF nao existe ou e' mais velho que o CSV."""
    pendentes = []
    for pasta in pastas:
        for n in arquivos(pasta):
            up = n.upper()
            if up.startswith(("LISTA_ATIVOS_SAUDE_", "LISTA_ATIVOS_ODONTO_")) and up.endswith(".CSV"):
                origem, destino = os.path.join(pasta, n), os.path.join(pasta, n[:-4] + ".pdf")
                if _desatualizado(origem, destino):
                    pendentes.append((origem, destino))
    if not pendentes:
        return 0
    gerados = 0
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        page = browser.new_page()
        try:
            for origem, destino in pendentes:
                try:
                    page.set_content(_html_lista(origem))
                    page.pdf(path=destino, format="A4", landscape=True, print_background=True,
                             margin={"top": "12mm", "bottom": "12mm", "left": "10mm", "right": "10mm"},
                             display_header_footer=True, header_template="<span></span>",
                             footer_template="<div style='font-size:8px;width:100%;text-align:center;color:#666'>"
                                             "Página <span class='pageNumber'></span> de <span class='totalPages'></span></div>")
                    gerados += 1
                except Exception as e:
                    log(f"pdf lista {os.path.basename(origem)}: ERRO {e}")
        finally:
            browser.close()
    return gerados


def _html_lista(caminho_csv: str) -> str:
    import html as _h
    with open(caminho_csv, "rb") as f:
        texto = _texto_censo(f.read())
    delim = ";" if texto.splitlines()[0].count(";") >= texto.splitlines()[0].count(",") else ","
    linhas = list(csv.reader(io.StringIO(texto), delimiter=delim, quotechar='"'))
    cab = [c.strip().upper() for c in linhas[0]]
    cols = []
    for rotulo, candidatos in _COLUNAS_LISTA_PDF:
        i = next((cab.index(c) for c in candidatos if c in cab), None)
        if i is not None:
            cols.append((rotulo, i))
    dados = [l for l in linhas[1:] if any(v.strip() for v in l)]
    # coluna sem nenhum valor (ex.: nascimento/admissao no Legado) nao entra no PDF
    cols = [(r, i) for r, i in cols if any(i < len(l) and l[i].strip() for l in dados)]
    i_tipo = next((i for r, i in cols if r == "Tipo"), None)
    titulares = sum(1 for l in dados if i_tipo is not None and i_tipo < len(l) and "TITULAR" in l[i_tipo].upper())
    nome_arq = os.path.basename(caminho_csv)[:-4]
    partes = nome_arq.split("_")  # LISTA_ATIVOS_{PLANO}_{COD}_{MES}
    plano, codigo, comp = partes[2], partes[3], partes[4]
    empresa = os.path.basename(os.path.dirname(os.path.dirname(os.path.dirname(caminho_csv))))
    th = "".join(f"<th>{_h.escape(r)}</th>" for r, _ in cols)
    corpo = "".join("<tr>" + "".join(f"<td>{_h.escape(l[i].strip() if i < len(l) else '')}</td>" for _, i in cols) + "</tr>"
                    for l in dados)
    resumo = f"{len(dados)} beneficiários" + (f" ({titulares} titulares, {len(dados) - titulares} dependentes)" if i_tipo is not None else "")
    return f"""<html><head><meta charset="utf-8"><style>
      body{{font-family:Arial,sans-serif;font-size:8px;color:#222}}
      h1{{font-size:14px;margin:0 0 2px}} p{{margin:0 0 8px;color:#555}}
      table{{border-collapse:collapse;width:100%}} th{{background:#1f3b5c;color:#fff;text-align:left;padding:3px 4px}}
      td{{border-bottom:1px solid #ddd;padding:1px 4px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;max-width:380px}}
      tr:nth-child(even) td{{background:#f6f6f6}}
      thead{{display:table-header-group}}
    </style></head><body>
      <h1>Lista de beneficiários ativos — {_h.escape(empresa)}</h1>
      <p>Contrato {codigo} · {"Saúde" if plano == "SAUDE" else "Odonto"} · competência {comp} · {resumo} ·
         fonte: relatório oficial da operadora</p>
      <table><thead><tr>{th}</tr></thead><tbody>{corpo}</tbody></table>
    </body></html>"""


_RE_DECIMAL = re.compile(r"^-?\d+\.\d{1,2}$")


def _csv_generico_para_xlsx(origem: str, destino: str) -> None:
    """Lista de ativos e demais CSVs: tudo como texto (CPF, matricula e
    codigos tem zero a esquerda), exceto valor com casas decimais (19.80)."""
    with open(origem, "rb") as f:
        texto = _texto_censo(f.read())
    amostra = texto.splitlines()[0] if texto else ""
    delim = ";" if amostra.count(";") >= amostra.count(",") else ","
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados"
    for i, linha in enumerate(csv.reader(io.StringIO(texto), delimiter=delim, quotechar='"')):
        valores = [ILLEGAL_CHARACTERS_RE.sub("", v) for v in linha]
        ws.append([float(v) if i and _RE_DECIMAL.match(v) else v for v in valores])
    for celula in ws[1]:
        celula.font = openpyxl.styles.Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save(destino)


def _csv_para_xlsx(origem: str, destino: str) -> None:
    with open(origem, encoding="utf-8-sig", errors="replace") as f:
        linhas = f.read().splitlines()
    i_cab = next(i for i, l in enumerate(linhas) if l.count(";") >= 5)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Analitico"
    for l in linhas[:i_cab]:
        if l.strip():
            ws.append([l.strip()])
    ws.append([])
    cab = linhas[i_cab].split(";")
    ws.append(cab)
    linha_cab = ws.max_row
    centavos = {i for i, c in enumerate(cab) if c.strip().lower() in COLUNAS_CENTAVOS}
    for l in linhas[i_cab + 1:]:
        if not l.strip():
            continue
        valores = l.split(";")
        ws.append([
            (_num(v) / 100 if isinstance(_num(v), int) else v) if i in centavos else v
            for i, v in enumerate(valores)
        ])
        for i in centavos:
            ws.cell(row=ws.max_row, column=i + 1).number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(row=linha_cab + 1, column=1)
    wb.save(destino)


# ---------------------------------------------------------------- auditoria

def codigos_ja_enviados(comp: str) -> set[str]:
    """Contratos cujo e-mail de faturamento da competencia ja' saiu (log do
    envio_faturamento). Nao baixa nem cobra arquivo deles: regra nova vale
    pra frente, sem reenvio (Andre, 01/10/2026)."""
    try:
        with open(os.path.join(BASE, "_status", f"ENVIOS_{comp}.json"), encoding="utf-8") as f:
            registros = json.load(f)
    except (OSError, ValueError):
        return set()
    return {c for r in registros if r.get("status") in ("enviado", "reenviado") for c in r["codigos"].split("+")}


def modalidades(codigos: set[str]) -> dict[str, str]:
    conn = psycopg2.connect(dbname="lifetools")
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT codigo_externo, modalidade FROM empresa_operadora_codigos "
                        "WHERE operadora_id = 1 AND codigo_externo = ANY(%s)", (list(codigos),))
            return dict(cur.fetchall())
    finally:
        conn.close()


def auditar(comp: str, saude: dict[str, set[str]], legado: dict[str, str], assumidos: set[str] = frozenset()) -> dict:
    """Padrao NOV26 (padrao_faturamento): cobra os nomes da lista fechada nos
    dois portais, confere total do relatorio x boleto do mes e aponta arquivo
    fora do padrao (que nao vai pro cliente)."""
    status_cw = status_legado(comp)
    anterior = competencia_anterior(comp)
    enviados = codigos_ja_enviados(comp)
    mods = modalidades(set(legado))
    linhas = []
    for codigo, tipos in sorted(saude.items()):
        pasta = pasta_contrato(codigo, comp)
        if codigo in enviados:
            linhas.append({"grupo": "Saúde (Code)", "codigo": codigo, "pasta": pasta, "falta": [], "obs": "já enviado"})
            continue
        falta = faltando_saude(codigo, tipos, comp, conferir=True)
        fora = pf.fora_do_padrao(arquivos(pasta))
        obs = f"fora do padrão (não vai pro cliente): {', '.join(fora)}" if fora else ""
        linhas.append({"grupo": "Saúde (Code)", "codigo": codigo, "pasta": pasta, "falta": falta, "obs": obs})
    for codigo, empresa in sorted(legado.items()):
        if codigo in saude:
            continue
        pasta = pasta_contrato(codigo, comp, empresa)
        grupo = "Legado (Code)" if codigo in assumidos else "Legado (Cowork)"
        if codigo in enviados:
            linhas.append({"grupo": grupo, "codigo": codigo, "pasta": pasta, "falta": [], "obs": "já enviado"})
            continue
        nomes = arquivos(pasta)
        st = status_cw.get(codigo)
        plano = pf.PLANO_DA_MODALIDADE.get(mods.get(codigo), "ODONTO")
        # Saude do Legado (Z8TUR, LH ERQ43/ICBJY): 2 boletos no mes = mensalidade + copart.
        copart = plano == "SAUDE" and bool(st) and len(st["boletos"]) >= 2
        falta = pf.faltando(nomes, codigo, comp, plano, copart)
        falta += [n for n in pf.faltando_fontes(nomes, codigo, comp, plano, copart) if not n.startswith("LISTA_ATIVOS")]
        obs = []
        if any(f.startswith("LISTA_ATIVOS") for f in falta) and tem(arquivos(pasta_contrato(codigo, anterior, empresa)), "LISTA_ATIVOS"):
            falta = [f for f in falta if not f.startswith("LISTA_ATIVOS")]
            obs.append(f"lista de ativos só em {anterior} (o envio copia se tiver até 10 dias)")
        nfs, nfs_outro_mes = nfs_do_mes(pasta, [n for n in nomes if _normaliza(n).startswith("NFE")], comp)
        falta += [f"NF de outro mês: {n}" for n in nfs_outro_mes]
        falta += pf.conferir_valores(pasta, codigo, comp, plano, copart, st["valores"] if st else set())
        if any(f.startswith("NFE_") for f in falta) and st and "NFE" not in st["tipos"]:
            obs.append(f"NF ainda não emitida pela Hapvida (status do {grupo[8:-1]} só tem boleto)")
        if falta and st and not st["emitido"]:
            obs.append(f"Hapvida ainda não emitiu nº de boleto/NF (status do {grupo[8:-1]})")
        fora = pf.fora_do_padrao(nomes)
        if fora:
            obs.append(f"fora do padrão (não vai pro cliente): {', '.join(fora)}")
        linhas.append({"grupo": grupo, "codigo": codigo, "pasta": pasta, "falta": falta, "obs": "; ".join(obs)})
    return {"competencia": comp, "gerado_em": datetime.now().isoformat(timespec="seconds"), "contratos": linhas}


def escrever_relatorio(aud: dict, modo: str, avisos: dict, erros_lista: list[str]) -> str:
    comp = aud["competencia"]
    completos = [l for l in aud["contratos"] if not l["falta"]]
    incompletos = [l for l in aud["contratos"] if l["falta"]]
    md = [
        f"# Auditoria de faturamento Hapvida — {comp}",
        "",
        f"Gerado em {aud['gerado_em']} (modo `{modo}`). "
        f"{len(completos)} de {len(aud['contratos'])} contratos completos.",
        "",
        "Padrão NOV26 (PROCEDIMENTO_MENSAL.md, seção \"Padrão de arquivos\"): por plano (Saúde, Coparticipação "
        "se houver, Odonto) boleto e NF em PDF, relatório analítico PDF+XLSX e lista de ativos XLSX+PDF, todos "
        "`{TIPO}_{PLANO}_{CÓDIGO}_{MÊS}`. O CSV oficial fica na pasta como fonte (XLSX/PDF gerados aqui) e o total "
        "do relatório tem que bater com o boleto do mês. CSV nunca vai ao cliente.",
        "",
        "## Faltando",
        "",
    ]
    if incompletos:
        md += ["| Grupo | Contrato | Falta | Observação |", "|---|---|---|---|"]
        for l in incompletos:
            md.append(f"| {l['grupo']} | {l['codigo']} | {', '.join(l['falta'])} | {l['obs']} |")
    else:
        md.append("Nada — todos os contratos completos.")
    md += ["", "## Completos", "", ", ".join(f"{l['codigo']}" for l in completos) or "—", ""]
    obs_completos = [l for l in completos if l["obs"]]
    if obs_completos:
        md += ["### Observações em contratos completos", ""] + [f"- {l['codigo']}: {l['obs']}" for l in obs_completos] + [""]
    avisos_restantes = {c: a for c, a in avisos.items() if a and any(l["codigo"] == c and l["falta"] for l in incompletos)}
    if avisos_restantes or erros_lista:
        md += ["## Avisos do download", ""]
        md += [f"- {c}: " + " | ".join(a) for c, a in sorted(avisos_restantes.items())]
        md += [f"- {e}" for e in erros_lista]
        md.append("")
    destino_dir = os.path.join(BASE, "_status")
    os.makedirs(destino_dir, exist_ok=True)
    caminho = os.path.join(destino_dir, f"AUDITORIA_FATURAMENTO_{comp}.md")
    with open(caminho, "w", encoding="utf-8") as f:
        f.write("\n".join(md))
    with open(os.path.join(destino_dir, f"AUDITORIA_FATURAMENTO_{comp}.json"), "w", encoding="utf-8") as f:
        json.dump(aud, f, ensure_ascii=False, indent=1)
    return caminho


def escrever_handoff_cowork(aud: dict) -> str | None:
    faltando = [l for l in aud["contratos"] if l["grupo"] == "Legado (Cowork)" and l["falta"]]
    if not faltando or not os.path.isdir(HANDOFF_ENTRADA):
        return None
    comp = aud["competencia"]
    corpo = [
        "DE: code | PARA: cowork",
        f"ASSUNTO: Faturamento {comp} — arquivos do Legado faltando na pasta (auditoria automática)",
        f"CONTEXTO: Auditoria automática do procedimento mensal (dia 26 00:01 completo; último dia do mês e dia 04 00:01 faltantes). Relatório completo em "
        f"`operadoras/hapvida/faturas/_status/AUDITORIA_FATURAMENTO_{comp}.md`.",
        "",
        "Contratos do Legado com arquivo faltando:",
    ]
    assumidos = sorted(l["codigo"] for l in aud["contratos"] if l["grupo"] == "Legado (Code)")
    if assumidos:
        corpo[-1:-1] = [f"O Code já baixa sozinho (login automático, credencial no Bitwarden) e NÃO precisa de você: "
                        f"{', '.join(assumidos)} (lista viva em `_status/LEGADO_ASSUMIDO_CODE.json`).", ""]
    corpo += [f"  - {l['codigo']}: {', '.join(l['falta'])}" + (f" ({l['obs']})" if l["obs"] else "") for l in faltando]
    corpo += [
        "",
        "PEDIDO: baixar o que falta na próxima execução do fluxo mensal, com os nomes do padrão NOV26 "
        "(operadoras/hapvida/PROCEDIMENTO_MENSAL.md, seção \"Padrão de arquivos\"). Deixe sempre o CSV oficial "
        "do relatório analítico (o Code confere o total com o boleto e gera o XLSX); se deixar o CSV da lista de "
        "ativos, o Code gera XLSX e PDF. Relatório com total diferente do boleto do mês = arquivo errado.",
        "ENTREGAR EM: pasta de faturas de sempre (sem resposta em _handoff/saida/ necessária)",
        "PRAZO/URGÊNCIA: normal",
    ]
    caminho = os.path.join(HANDOFF_ENTRADA, f"{date.today():%Y%m%d}_code_para_cowork_faturamento_{comp.lower()}_faltando.md")
    with open(caminho, "w", encoding="utf-8") as f:
        f.write("\n".join(corpo) + "\n")
    return caminho


# ---------------------------------------------------------------- legado (Code)

def legado_assumido(legado: dict[str, str]) -> tuple[set[str], dict]:
    """Contratos do Legado que o Code baixa sozinho (legado_auto): Odonto e
    Saude (08/10/2026) com credencial no segredo hapvida-legado (copiado do Bitwarden pelo
    sync_bitwarden_legado.py). Grava a lista em _status/LEGADO_ASSUMIDO_CODE.json
    pro Cowork pular esses."""
    try:
        creds = secrets_store.get_json("hapvida-legado")
    except Exception:
        creds = {}
    mods = modalidades(set(legado))
    # CNPJ sem 14 digitos = Troca de Arquivos vai recusar (Y4VLX 08/10): fica com o Cowork ate corrigir
    assumidos = {c for c in legado if c in creds and pf.PLANO_DA_MODALIDADE.get(mods.get(c))
                 and len(creds[c].get("cnpj") or "") == 14}
    destino = os.path.join(BASE, "_status", "LEGADO_ASSUMIDO_CODE.json")
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    with open(destino, "w", encoding="utf-8") as f:
        json.dump({"atualizado_em": datetime.now().isoformat(timespec="seconds"), "codigos": sorted(assumidos),
                   "obs": "Contratos do Legado baixados pelo Code (legado_auto.py). O Cowork não baixa estes."},
                  f, ensure_ascii=False, indent=1)
    return assumidos, creds


def baixar_legado(legado: dict[str, str], assumidos: set[str], creds: dict, comp: str) -> dict[str, list[str]]:
    mes, ano = competencia_para_mes(comp)
    mods = modalidades(assumidos)
    enviados = codigos_ja_enviados(comp)
    avisos: dict[str, list[str]] = {}
    status: list[dict] = []
    for codigo in sorted(assumidos - enviados):
        try:
            r = legado_auto.baixar_contrato(codigo, legado[codigo], creds[codigo], comp, mes, ano,
                                            pasta_contrato(codigo, comp, legado[codigo]),
                                            pf.PLANO_DA_MODALIDADE[mods[codigo]])
            avisos[codigo] = r["avisos"]
            status += r["status"]
        except Exception as e:
            avisos[codigo] = [f"Legado: {e}"]
            log(f"legado {codigo}: ERRO {e}")
    legado_auto.gravar_status(BASE, comp, status)
    return avisos


# ---------------------------------------------------------------- main

def atualizar_faturas(creds: dict) -> list[str]:
    """Faturas em `faturas` (inclui unpaid-bills, correcao 28/09) - senao a
    competencia nova nao aparece e o download pega o mes anterior."""
    # 02/10/2026: falha aqui (timeout da Hapvida no DAVJV em 01/10) derrubava a
    # rodada inteira antes do download e das listas. Agora cada parte loga o
    # erro e segue - o que nao atualizou fica com o que ja estava no banco.
    erros = []
    try:
        r = boletos_saude.run_boletos_saude(creds["carol"]["usuario"], creds["carol"]["senha"])
        falhas = [x["codigo"] for x in r["resultados"] if x["status"] != "ok"]
        log(f"faturas (Carol): {r['total_novos']} novas, {r['total_atualizados']} atualizadas"
            + (f", {len(falhas)} código(s) com erro: {', '.join(falhas)}" if falhas else ""))
    except Exception as e:
        erros.append(f"Atualizar faturas (Carol): {e}")
        log(f"faturas (Carol): ERRO {e}")
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            page = browser.new_page()
            try:
                auth = login_saude(page, creds["marco"]["usuario"], creds["marco"]["senha"])
            finally:
                browser.close()
        fim = date.today().replace(year=date.today().year + 1).isoformat()
        linhas = boletos_saude.capturar_boletos_contrato(auth, "DAVJV", "2025-01-01", fim)
        r = boletos_saude.upsert_faturas(linhas, fonte=f"davjv_marco_{date.today().isoformat()}")
        log(f"faturas DAVJV (Marcos): {r['novos']} novas, {r['atualizados']} atualizadas")
    except Exception as e:
        erros.append(f"Atualizar faturas DAVJV (Marcos): {e}")
        log(f"faturas DAVJV (Marcos): ERRO {e}")
    return erros


def main() -> None:
    args = dict(a.lstrip("-").split("=", 1) for a in sys.argv[1:] if "=" in a)
    # --modo obrigatorio (05/10/2026: um '--help' rodou o procedimento de faltantes sem querer).
    modo = args.get("modo")
    if modo not in ("completo", "faltantes"):
        sys.exit("uso: faturamento_mensal.py --modo=completo|faltantes [--competencia=OUT26]")
    comp = args.get("competencia") or competencia_padrao(date.today())
    so_auditoria = "--so-auditoria" in sys.argv
    if not so_auditoria and os.path.exists(os.path.join(BASE, "_status", "FATURAMENTO_PAUSADO")):
        log("FATURAMENTO_PAUSADO presente em _status/ - nada baixado (só auditoria).")
        so_auditoria = True
    log(f"inicio - modo {modo}, competencia {comp}")

    project = os.environ["GCP_PROJECT"]
    avisos: dict = {}
    erros_lista: list[str] = []
    saude: dict = {}
    if not so_auditoria:
        creds = {
            "carol": scraper.get_credentials(project, "hapvida-portal-saude"),
            "marco": scraper.get_credentials(project, "hapvida-portal-saude-marco"),
        }
        erros_lista += atualizar_faturas(creds)
        saude = contratos_saude(comp)
        enviados = codigos_ja_enviados(comp)
        a_baixar = {c: t for c, t in saude.items() if c not in enviados}
        log(f"{len(saude)} contrato(s) Saúde com mensalidade em {comp}, {len(saude) - len(a_baixar)} já enviado(s)")
        try:
            avisos = baixar_documentos_saude(a_baixar, comp, modo, creds)
        except Exception as e:
            erros_lista.append(f"Download de documentos Saúde: {e}")
            log(f"download Saúde: ERRO {e}")
        try:
            erros_lista += salvar_listas_ativos(a_baixar, comp, modo, creds)
        except Exception as e:
            erros_lista.append(f"Listas de ativos: {e}")
            log(f"listas de ativos: ERRO {e}")
    else:
        saude = contratos_saude(comp)

    legado = contratos_legado()
    assumidos, creds_legado = legado_assumido(legado)
    if not so_auditoria and assumidos:
        log(f"Legado pelo Code: {', '.join(sorted(assumidos))}")
        avisos.update(baixar_legado(legado, assumidos, creds_legado, comp))

    log(f"xlsx gerados: {gerar_xlsx_relatorios(comp)}")
    aud = auditar(comp, saude, legado, assumidos)
    caminho = escrever_relatorio(aud, modo, avisos, erros_lista)
    handoff = escrever_handoff_cowork(aud)
    incompletos = [l["codigo"] for l in aud["contratos"] if l["falta"]]
    log(f"auditoria: {caminho} | incompletos: {incompletos or 'nenhum'} | handoff: {handoff or '-'}")


if __name__ == "__main__":
    main()
